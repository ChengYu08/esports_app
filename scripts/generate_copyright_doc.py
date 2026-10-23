from __future__ import annotations

import textwrap
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOFTWARE_NAME = "盼之电竞移动客户端"
SOFTWARE_VERSION = "V1.0.0"
SOFTWARE_TITLE = f"{SOFTWARE_NAME} {SOFTWARE_VERSION} 源程序"
OUTPUT_FILE = Path("docs/软件著作权-盼之电竞移动客户端-源程序60页.docx")

SOURCE_ROOTS = ["src", "modules", "packages", "scripts"]
SOURCE_EXTENSIONS = {".ts", ".tsx", ".js", ".jsx", ".cjs", ".mjs"}
EXCLUDED_PARTS = {"node_modules", ".expo", "__tests__", "coverage", "dist", "build"}
CODE_PAGE_LINES = 65
FIRST_PAGE_CODE_LINES = 60
TOTAL_PAGES = 60
CODE_LINE_WIDTH = 100


@dataclass(frozen=True)
class CodeFile:
    path: Path
    lines: list[str]


def should_include_file(path: Path) -> bool:
    if path.suffix.lower() not in SOURCE_EXTENSIONS:
        return False
    return not any(part in EXCLUDED_PARTS for part in path.parts)


def discover_code_files(project_root: Path) -> list[Path]:
    files: list[Path] = []
    for root_name in SOURCE_ROOTS:
        root = project_root / root_name
        if not root.exists():
            continue
        files.extend(
            path
            for path in root.rglob("*")
            if path.is_file() and should_include_file(path)
        )
    files.sort(key=lambda path: path.as_posix())
    return files


def load_code_files(project_root: Path, paths: list[Path]) -> list[CodeFile]:
    code_files: list[CodeFile] = []
    for path in paths:
        relative = path.relative_to(project_root).as_posix()
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="utf-8", errors="replace")
        lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        code_files.append(CodeFile(path=relative, lines=lines))
    return code_files


def wrap_code_line(line: str) -> list[str]:
    line = line.rstrip("\n").replace("\t", "    ")
    if not line:
        return [""]
    return textwrap.wrap(
        line,
        width=CODE_LINE_WIDTH,
        break_long_words=True,
        break_on_hyphens=False,
        subsequent_indent="    ",
    )


def build_code_lines(code_files: list[CodeFile]) -> list[str]:
    lines: list[str] = []
    for code_file in code_files:
        if lines:
            lines.append("")
        lines.extend(wrap_code_line(f"// ===== File: {code_file.path} ====="))
        lines.append("")
        for line in code_file.lines:
            lines.extend(wrap_code_line(line))
    return lines


def set_font(run, size: float, bold: bool = False, east_asia: str = "宋体") -> None:
    run.font.name = "Consolas"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_code_line(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph_format.line_spacing = Pt(10.8)
    run = paragraph.add_run(text)
    set_font(run, 8.5)


def add_title_page_header(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run(SOFTWARE_TITLE)
    set_font(title_run, 15, bold=True, east_asia="黑体")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle_run = subtitle.add_run("源程序（前 30 页与后 30 页，中间部分略）")
    set_font(subtitle_run, 10, east_asia="宋体")


def add_page_break(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_page_number_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    set_font(run, 8.5, east_asia="宋体")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_end)


def configure_page(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.7)


def configure_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_run = header_paragraph.add_run(f"{SOFTWARE_NAME} {SOFTWARE_VERSION} 源程序")
    set_font(header_run, 8.5, east_asia="宋体")

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    prefix = footer_paragraph.add_run("第 ")
    set_font(prefix, 8.5, east_asia="宋体")
    add_page_number_field(footer_paragraph, "PAGE")
    middle = footer_paragraph.add_run(" 页 / 共 ")
    set_font(middle, 8.5, east_asia="宋体")
    add_page_number_field(footer_paragraph, "NUMPAGES")
    suffix = footer_paragraph.add_run(" 页")
    set_font(suffix, 8.5, east_asia="宋体")


def write_code_pages(document: Document, code_lines: list[str]) -> None:
    if len(code_lines) < TOTAL_PAGES * CODE_PAGE_LINES:
        raise ValueError(
            "源码行数不足，无法生成前 30 页与后 30 页。"
            f"当前可排版行数：{len(code_lines)}"
        )

    first_page_line_count = FIRST_PAGE_CODE_LINES
    front_line_count = first_page_line_count + (TOTAL_PAGES // 2 - 1) * CODE_PAGE_LINES
    back_line_count = CODE_PAGE_LINES * (TOTAL_PAGES // 2)

    front_lines = code_lines[:front_line_count]
    back_lines = code_lines[-back_line_count:]

    add_title_page_header(document)
    for index, line in enumerate(front_lines, start=1):
        add_code_line(document, line)
        if index == first_page_line_count:
            break
    add_page_break(document)

    page_index = 2
    remaining_front_lines = front_lines[first_page_line_count:]
    while remaining_front_lines:
        page_lines = remaining_front_lines[:CODE_PAGE_LINES]
        remaining_front_lines = remaining_front_lines[CODE_PAGE_LINES:]
        for line in page_lines:
            add_code_line(document, line)
        add_page_break(document)
        page_index += 1

    while page_index < TOTAL_PAGES:
        start = (page_index - 31) * CODE_PAGE_LINES
        page_lines = back_lines[start : start + CODE_PAGE_LINES]
        for line in page_lines:
            add_code_line(document, line)
        add_page_break(document)
        page_index += 1

    last_page_start = (TOTAL_PAGES - 31) * CODE_PAGE_LINES
    for line in back_lines[last_page_start:]:
        add_code_line(document, line)


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    output_path = project_root / OUTPUT_FILE
    output_path.parent.mkdir(parents=True, exist_ok=True)

    code_files = load_code_files(project_root, discover_code_files(project_root))
    code_lines = build_code_lines(code_files)

    document = Document()
    section = document.sections[0]
    configure_page(section)
    configure_header_footer(section)
    write_code_pages(document, code_lines)

    document.save(output_path)
    print(output_path.as_posix())
    print(f"files={len(code_files)}")
    print(f"lines={len(code_lines)}")
    print(f"pages={TOTAL_PAGES}")


if __name__ == "__main__":
    main()
